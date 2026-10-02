# -*- coding: utf-8 -*-
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
import shutil

def create_report():
    doc = Document()

    # 1. 页面边距设置 (上/下 2cm，左/右 1.8cm，提供更充裕的表格横向展示空间)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    # 基础字体与段落工具函数
    def set_run_font(run, font_name="微软雅黑", size_pt=10, bold=False, color_rgb=None):
        run.font.name = font_name
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
        run.font.size = Pt(size_pt)
        run.bold = bold
        if color_rgb:
            run.font.color.rgb = color_rgb

    def add_custom_heading(text, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        if level == 1:
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.2
            run = p.add_run(text)
            set_run_font(run, "微软雅黑", 13.5, bold=True, color_rgb=RGBColor(31, 78, 121)) # 深商务蓝
        return p

    def add_body_p(text="", bold_prefix="", space_after=4, line_spacing=1.25):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            set_run_font(r_pre, "微软雅黑", 10, bold=True, color_rgb=RGBColor(38, 38, 38))
        if text:
            r_txt = p.add_run(text)
            set_run_font(r_txt, "微软雅黑", 10, bold=False, color_rgb=RGBColor(60, 60, 60))
        return p

    def set_cell_shading(cell, color_hex):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
                <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideV w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr.append(borders)

    # ---------------- 标题部分 ----------------
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(4)
    title_p.paragraph_format.space_after = Pt(3)
    run_title = title_p.add_run("关于简历自动填写插件在 10 家银行网申系统的实测情况汇报")
    set_run_font(run_title, "微软雅黑", 17, bold=True, color_rgb=RGBColor(31, 78, 121))

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(12)
    run_sub = sub_p.add_run("测试范围：工商银行、农业银行、建设银行、中国银行、交通银行、邮储银行、招商银行、浦发银行、中信银行、兴业银行")
    set_run_font(run_sub, "微软雅黑", 9, color_rgb=RGBColor(120, 120, 120))

    # ---------------- 一、 测试概况 ----------------
    add_custom_heading("一、 测试概况", level=1)
    add_body_p("本次测试主要检验了“简历自动填写插件”在 10 家国内主流商业银行招聘系统（网申个人简历填报页面）的实际表现。")
    add_body_p("从整体测试情况来看，大部分银行的基本信息（如个人资料、教育背景、工作与实习经历等）都能比较顺畅地自动填入。插件表现最好的银行基本在 6～9 秒内就能填完大半；主要的卡点和失分项集中在：部分下拉菜单选不上、分步骤弹窗识别不灵、个别字段填错位置，以及招商银行页面目前完全不响应。")

    # ---------------- 二、 10 家银行实测表现全景汇总表 ----------------
    add_custom_heading("二、 10 家银行实测表现全景汇总表", level=1)

    table_data = [
        ["序号", "银行名称", "页面形式", "填写耗时", "填写完整度\n(大致估算)", "实际测试情况与具体表现详述"],
        [
            "1", "建设银行", "一个单页面", "6 秒", "接近 100%",
            "整体表现最为顺畅，基础信息、教育经历、各项经历全部精准自动填入。仅个别银行内部专属关联选项（如“是否在建行工作”）未填，基本不需要人工再补内容。"
        ],
        [
            "2", "交通银行", "一个单页面", "6 秒", "90% 左右",
            "日常必要字段基本全部填入，填报非常完整顺畅。填写过程中偶尔会弹出一个提示小窗口，随手关掉即可，不影响实际填充功能；亲属关系等简历中没有的信息正常留空。"
        ],
        [
            "3", "中信银行", "一个单页面", "8 秒", "90% 左右",
            "主体绝大部分内容都填得很好，速度也快。主要问题集中在局部：① 薪酬福利那一整块漏填；② 教育经历里的“毕业方式”未填；③ 工作经历因是下拉框形式未能选入；④ 其他重要信息中个别难以推断的选择项未选。"
        ],
        [
            "4", "兴业银行", "一个单页面", "9 秒", "85% 左右",
            "整体填入速度很快。主要问题包括：① 籍贯、户口所在地等带有查询搜索性质的下拉框，插件直接填入了系统数据源的一整串长文本；② 求职地点受系统选项限制未能选上；③ 专业证书发证机构漏填；④ 竞业限制等合规声明因简历无源数据留空。"
        ],
        [
            "5", "中国银行", "一个单页面", "18 秒", "80% 左右",
            "常规字段填得都比较规整完整。耗时比其他单页银行稍长（18秒）；“是否具有境外身份”没选上；像“受雇于中行的亲属”这类特殊背景调查字段正常留空，其余主体内容覆盖良好。"
        ],
        [
            "6", "工商银行", "模块化分块", "各模块 1~2 秒", "80% 左右",
            "大多数大块内容均能正常填入。但页面割裂导致操作较繁琐，需要用户手动逐个点开卡片触发插件；部分下拉框（如语言水平）选得不太好；像服从调剂、获取招聘渠道等非必填项没有填。"
        ],
        [
            "7", "浦发银行", "多个页面分开", "需手动逐页点", "80% 左右",
            "填入的大部分内容算正常。主要问题：① 交互割裂，需手动频繁打开多个独立页面分次点击；② 最开始的国家、地区、省市级联下拉框未选上；③ 存在字段错位现象，“自我评价”一栏被填成了其他的杂项说明信息。"
        ],
        [
            "8", "中国邮政", "标签页手动切", "需手动切页点", "80% 左右",
            "主干的实习经历、工作经历和基本资料填得挺好。主要问题：① 存在严重填串位置的现象（把姓名框填成了手机号码）；② 科研经历、出版物、专利及专业资格证书未能填入；③ “个人爱好”输入框未被插件检出。"
        ],
        [
            "9", "农业银行", "弹窗逐个添加", "部分弹窗 0 秒", "80% 左右",
            "操作体验比较繁琐，必须逐个点“添加”在浮层弹窗内填写。主要问题：① 弹窗识别极不稳定，部分弹窗能顺畅填完，但有些弹窗插件压根检测不到（直接显示0秒）；② 里面的下拉选项和圆圈单选框大多没选上；③ 出现严重错配，在“技能和证书”弹窗里填入了“学历”内容。"
        ],
        [
            "10", "招商银行", "一个单页面", "0 秒", "0% (不可用)",
            "不论点击“一键填入”还是“选区填入”，插件均直接显示耗时 0 秒，页面完全没有任何反应，所有输入框都填不进去，目前处于完全不兼容状态。"
        ]
    ]

    table = doc.add_table(rows=len(table_data), cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="D3D3D3", sz="4")

    # 列宽比例设定（总宽约 6.87 英寸）
    # 序号 0.4 | 银行 0.85 | 页面 0.9 | 耗时 0.85 | 完整度 0.95 | 详细说明 2.92
    col_widths = [Inches(0.4), Inches(0.85), Inches(0.9), Inches(0.85), Inches(0.95), Inches(2.92)]

    for row_idx, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths[col_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=90, bottom=90, left=100, right=100)

            text = table_data[row_idx][col_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.2

            # 对齐方式
            if col_idx in [0, 1, 2, 3, 4]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            if row_idx == 0:
                # 表头
                set_cell_shading(cell, "1F4E79") # 商务深蓝
                run = p.add_run(text)
                set_run_font(run, "微软雅黑", 9.5, bold=True, color_rgb=RGBColor(255, 255, 255))
            else:
                # 数据行
                bg_color = "F9FAFB" if row_idx % 2 == 0 else "FFFFFF"
                set_cell_shading(cell, bg_color)
                run = p.add_run(text)
                
                # 强调颜色
                if col_idx == 4: # 完整度
                    if "100%" in text or "90%" in text:
                        set_run_font(run, "微软雅黑", 9.5, bold=True, color_rgb=RGBColor(46, 125, 50)) # 绿色
                    elif "0%" in text:
                        set_run_font(run, "微软雅黑", 9.5, bold=True, color_rgb=RGBColor(198, 40, 40)) # 红色
                    else:
                        set_run_font(run, "微软雅黑", 9.5, bold=True, color_rgb=RGBColor(31, 78, 121))
                elif col_idx == 1:
                    set_run_font(run, "微软雅黑", 9.5, bold=True, color_rgb=RGBColor(38, 38, 38))
                elif col_idx == 5:
                    set_run_font(run, "微软雅黑", 8.8, bold=False, color_rgb=RGBColor(55, 55, 55))
                else:
                    set_run_font(run, "微软雅黑", 9, bold=False, color_rgb=RGBColor(70, 70, 70))

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ---------------- 三、 测试中发现的几类常见问题 ----------------
    add_custom_heading("三、 测试中发现的几类常见问题（通俗梳理）", level=1)
    add_body_p("从这次实测来看，插件在各家银行的表现主要受以下几个方面影响：")

    issues = [
        ("1. 单页填报比多页/弹窗填报体验好太多：", "像建行、交行、中信、中行、兴业这种把所有信息放一个大页面里的，点一次按键几秒钟就填完一大半，体验非常流畅；而像工行、浦发、邮政、农行，简历被拆成了好几页或者一个个弹窗，用户得手动点开一个、点一下插件、保存，再点下一个，稍微有些繁琐。"),
        ("2. 部分下拉菜单和选框容易漏选：", "很多银行的下拉框不是普通的简单下拉，而是带联动的（比如选了国家再选省份）、或者带搜索框的。比如浦发的省市下拉、中信的工作经历下拉、工行的语言能力，插件容易选不上；兴业银行的籍贯搜索框，直接把一串系统文字一股脑填进去了；农行的圆圈单选题，插件也经常勾不上。"),
        ("3. 少数位置出现“填错地方”的情况（需人工复核）：", "邮政把姓名填成了手机号；农行在技能证书里填了学历；浦发在自我评价里填了其他说明。这类情况不多，但属于明显填错，后续需要人工核对。"),
        ("4. 正常留空与确实漏填的区别：", "像“有没有亲属在银行上班”、“在银行有没有熟人”、“有无违法乱纪”、“境外身份”等，在候选人简历里本来就不会写，插件没填是完全合情合理的；而像中信的薪酬期望和毕业方式、兴业的证书发证机构等，则属于插件还可以进一步完善的地方。")
    ]

    for title, desc in issues:
        add_body_p(desc, bold_prefix=title, space_after=6)

    # 保存 Word 文档
    out_path = r"d:\100_Work\101_Program\Proj\XHS\关于简历自动填写插件在10家银行网申系统的实测情况汇报.docx"
    doc.save(out_path)
    shutil.copyfile(out_path, r"d:\100_Work\101_Program\Proj\XHS\bank_autofill_test_report.docx")
    print(f"Success! Document saved to {out_path} and bank_autofill_test_report.docx")

if __name__ == "__main__":
    create_report()
